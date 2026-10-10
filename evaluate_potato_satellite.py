"""Reproduce a small, field-held-out Sentinel-2 potato yield experiment.

The workbook is Al-Gaadi et al. (2016), PLOS ONE, S1 Table (CC BY 4.0):
https://plos.figshare.com/articles/dataset/3818955
It covers three Saudi potato fields in one growing season, not Pakistan.
"""

from __future__ import annotations

import csv
import json
import statistics
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / ".vendor"))
from openpyxl import load_workbook  # noqa: E402

SOURCE = ROOT / "data" / "sources" / "potato_paper" / "S1_Table.xlsx"
OUT = ROOT / "outputs"
FIELDS = {
    "67-S": ("Pivot 67-S Model", 14),
    "68-S": ("Pivot 68-S Model ", 10),
    "44-S": ("Pivot 44-S Model", 14),
}
SOURCE_URL = "https://plos.figshare.com/articles/dataset/3818955"
PAPER_URL = "https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0162219"
ALPHA = 1.0  # Identical in every fold; see README on exploratory selection.


def source_rows():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Published potato table missing: {SOURCE}")
    book = load_workbook(SOURCE, read_only=True, data_only=True)
    records = []
    for field, (sheet_name, ndvi_col) in FIELDS.items():
        sheet = book[sheet_name]
        for original_row, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
            # Read only the original *Model* sheets. In the published Validation
            # sheets, several VI columns contain fitted yield values, not raw VI.
            if not isinstance(row[16], (int, float)):
                continue
            if not isinstance(row[ndvi_col], (int, float)) or not 0 <= row[ndvi_col] <= 1:
                raise ValueError(f"Invalid Sentinel NDVI in {sheet_name}, row {original_row}")
            records.append({
                "field": field,
                "sample_id": str(row[0]),
                "source_row": original_row,
                "sentinel_ndvi": float(row[ndvi_col]),
                "measured_yield_t_ha": float(row[16]),
            })
    counts = {field: sum(row["field"] == field for row in records) for field in FIELDS}
    if counts != {field: 45 for field in FIELDS}:
        raise ValueError(f"Expected 45 original measured points per field; got {counts}")
    return records


def fit_one_feature_ridge(training):
    """Training-only standardization, one real satellite predictor, fixed ridge."""
    x = [row["sentinel_ndvi"] for row in training]
    y = [row["measured_yield_t_ha"] for row in training]
    x_mean = statistics.mean(x)
    x_std = statistics.pstdev(x)
    y_mean = statistics.mean(y)
    if x_std == 0:
        return lambda ndvi: y_mean
    normalized = [(value - x_mean) / x_std for value in x]
    slope = sum(value * (target - y_mean) for value, target in zip(normalized, y)) / (
        sum(value * value for value in normalized) + ALPHA
    )
    return lambda ndvi: y_mean + slope * ((ndvi - x_mean) / x_std)


def build_result(write_files=True):
    records = source_rows()
    field_scores = []
    prediction_rows = []
    for held_out in FIELDS:
        training = [row for row in records if row["field"] != held_out]
        test = [row for row in records if row["field"] == held_out]
        model = fit_one_feature_ridge(training)
        baseline = statistics.mean(row["measured_yield_t_ha"] for row in training)
        for row in test:
            prediction_rows.append({
                **row,
                "predicted_yield_t_ha": round(model(row["sentinel_ndvi"]), 3),
                "baseline_yield_t_ha": round(baseline, 3),
            })
        subset = [row for row in prediction_rows if row["field"] == held_out]
        model_mae = statistics.mean(abs(row["predicted_yield_t_ha"] - row["measured_yield_t_ha"]) for row in subset)
        baseline_mae = statistics.mean(abs(row["baseline_yield_t_ha"] - row["measured_yield_t_ha"]) for row in subset)
        field_scores.append({
            "field": held_out,
            "train_samples": len(training),
            "held_out_samples": len(test),
            "model_mae_t_ha": round(model_mae, 3),
            "baseline_mae_t_ha": round(baseline_mae, 3),
            "model_beats_baseline": model_mae < baseline_mae,
        })
    result = {
        "crop": "Potato",
        "status": "historical research demonstration",
        "place": "Wadi Al-Dawasir, Saudi Arabia",
        "season": "2016",
        "satellite": "Sentinel-2",
        "feature": "February 11, 2016 NDVI at published sample locations",
        "outcome": "Measured potato yield from harvested 3 m² sample, scaled to t/ha",
        "validation": "Leave one entire field out; train on the other two fields; same ridge penalty 1.0 in each fold; compare with training-field mean. Exploratory model choice needs independent confirmation.",
        "sample_count": len(prediction_rows),
        "field_count": len(FIELDS),
        "model_mae_t_ha": round(statistics.mean(abs(row["predicted_yield_t_ha"] - row["measured_yield_t_ha"]) for row in prediction_rows), 3),
        "baseline_mae_t_ha": round(statistics.mean(abs(row["baseline_yield_t_ha"] - row["measured_yield_t_ha"]) for row in prediction_rows), 3),
        "field_scores": field_scores,
        "points": prediction_rows,
        "source_url": SOURCE_URL,
        "paper_url": PAPER_URL,
        "scope_note": "Three nearby fields in one Saudi season are too little evidence for a reliable current forecast, Pakistan transfer, or supply/dispatch decision. One held-out field loses to the simple baseline. Model choices were explored on this dataset, so the measured advantage needs independent confirmation. The displayed points are 2016 held-out evaluations, not live predictions.",
    }
    if write_files:
        OUT.mkdir(exist_ok=True)
        (OUT / "potato_satellite_evaluation.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        with (OUT / "potato_satellite_points.csv").open("w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=list(prediction_rows[0]))
            writer.writeheader()
            writer.writerows(prediction_rows)
    return result


if __name__ == "__main__":
    result = build_result()
    print(json.dumps({key: value for key, value in result.items() if key not in ("points",)}, indent=2))
